"""Regression tests for additive viewing upgrades and shared family outcomes."""
import importlib
import io
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from tests.test_multiuser import multi_app
from tests.test_viewings import current, payload, create


def test_legacy_viewings_migrate_without_rebuild_or_loss(multi_app):
    main,admin=multi_app
    item=create(admin,status='completed',outcome='thinking',preparation='До поездки',comment='Старый общий итог')
    with main.database() as conn:
        conn.execute('DROP INDEX property_visits_request')
        for field in ('rejection_reasons','request_id','request_hash'):
            conn.execute('ALTER TABLE property_visits DROP COLUMN '+field)
        conn.execute('ALTER TABLE property_visits ADD COLUMN vendor_note TEXT')
        conn.execute("UPDATE property_visits SET vendor_note='Сохраняется'")
        conn.execute('CREATE INDEX vendor_visits_idx ON property_visits(vendor_note)')
        conn.execute("CREATE TRIGGER vendor_visits_trigger AFTER UPDATE OF vendor_note ON property_visits BEGIN SELECT 1; END")
        visits=[dict(row) for row in conn.execute('SELECT * FROM property_visits')]
        listings=[dict(row) for row in conn.execute('SELECT * FROM listings ORDER BY id')]
        schema={r['name']:r['sql'] for r in conn.execute("SELECT name,sql FROM sqlite_master WHERE name LIKE 'vendor_%'")}
    from app.migrations import migrate_listings
    migrate_listings(main.database);migrate_listings(main.database)
    with main.database() as conn:
        after=[dict(row) for row in conn.execute('SELECT * FROM property_visits')]
        for before,row in zip(visits,after):
            assert {k:row[k] for k in before}==before
            assert row['rejection_reasons']=='[]'
            assert row['request_id'] is None and row['request_hash'] is None
        assert [dict(row) for row in conn.execute('SELECT * FROM listings ORDER BY id')]==listings
        assert {r['name']:r['sql'] for r in conn.execute("SELECT name,sql FROM sqlite_master WHERE name LIKE 'vendor_%'")}==schema
        assert conn.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert conn.execute('PRAGMA foreign_key_check').fetchone() is None
    assert admin.get('/api/auth/me').status_code==200
    assert current(admin)['visits'][0]['outcome']=='thinking'


def test_both_notes_reasons_actual_date_reload_and_export(multi_app):
    main,admin=multi_app
    item=create(admin,status='completed',outcome='rejected',scheduled_at='2026-10-20T18:00:00+03:00',
                completed_at='2026-10-21T19:30:00+03:00',preparation='Что спросить до поездки',
                comment='=Без исполнения формул',rejection_reasons=['noise','documents'])
    visit=item['visits'][0]
    assert visit['completed_at']=='2026-10-21T16:30:00+00:00'
    assert visit['scheduled_at']=='2026-10-20T15:00:00+00:00'
    assert visit['rejection_reasons']==['noise','documents']
    assert 'request_id' not in visit and 'request_hash' not in visit
    assert admin.post('/api/listings/123/review',json={'status':'liked'}).status_code==200
    data=payload(item,status='completed',outcome='rejected',version=visit['version'],
                 completed_at=visit['completed_at'],preparation='Уточнённый вопрос',
                 comment='=После поездки',rejection_reasons=['layout','other'])
    response=admin.patch('/api/listings/123/visits/'+str(visit['id']),json=data)
    assert response.status_code==200 and response.json()['review_status']=='liked'
    importlib.reload(main)
    saved=current(admin)['visits'][0]
    assert saved['preparation']=='Уточнённый вопрос' and saved['comment']=='=После поездки'
    assert saved['rejection_reasons']==['layout','other'] and saved['completed_at']==visit['completed_at']
    book=load_workbook(io.BytesIO(admin.get('/api/export.xlsx').content))
    sheet=book['Просмотры']
    assert sheet.cell(2,8).value=='Уточнённый вопрос'
    assert sheet.cell(2,9).value=="'=После поездки"
    assert sheet.cell(1,13).value=='Причины отказа'
    assert sheet.cell(2,13).value=='Планировка, Другое'


@pytest.mark.parametrize('changes',[
    {'rejection_reasons':['unknown']},{'rejection_reasons':['noise','noise']},
    {'rejection_reasons':['noise']*8},{'rejection_reasons':['noise'],'outcome':'liked'},
    {'status':'planned','outcome':None,'rejection_reasons':['noise']},
    {'completed_at':'2026-10-20T19:30:00'},
    {'status':'planned','outcome':None,'completed_at':'2026-10-20T19:30:00+03:00'},
    {'preparation':'x'*2001},{'comment':'x'*2001},{'request_id':'not-a-uuid'},
])
def test_extended_visit_input_validation(multi_app,changes):
    main,admin=multi_app
    data=payload(current(admin),status='completed',outcome='rejected',**{})
    data.update(changes)
    assert admin.post('/api/listings/123/visits',json=data).status_code==422
    assert current(admin)['visits']==[]


def test_reasons_are_optional_preserved_for_legacy_clients_and_clear_on_positive(multi_app):
    main,admin=multi_app
    item=create(admin,status='completed',outcome='rejected',rejection_reasons=['noise'])
    visit=item['visits'][0];url='/api/listings/123/visits/'+str(visit['id'])
    response=admin.patch(url,json=payload(item,status='completed',outcome='rejected',version=visit['version'],comment='Только комментарий'))
    assert response.status_code==200
    visit=response.json()['visits'][0]
    assert visit['rejection_reasons']==['noise'] and visit['completed_at']
    response=admin.patch(url,json=payload(item,status='completed',outcome='liked',version=visit['version']))
    assert response.status_code==200 and response.json()['visits'][0]['rejection_reasons']==[]
    assert response.json()['review_status']=='liked'


def test_actual_chronology_not_creation_order_controls_family_outcome(multi_app):
    main,admin=multi_app
    newer=create(admin,status='completed',outcome='liked',completed_at='2026-10-25T18:00:00+03:00')
    older=create(admin,status='completed',outcome='rejected',completed_at='2026-10-20T18:00:00+03:00',rejection_reasons=['noise'])
    assert older['review_status']=='liked'
    old=older['visits'][0]
    response=admin.patch('/api/listings/123/visits/'+str(old['id']),json=payload(older,
        version=old['version'],status='completed',outcome='thinking',completed_at=old['completed_at']))
    assert response.status_code==200 and response.json()['review_status']=='liked'
    assert response.json()['visits'][0]['rejection_reasons']==[]


def test_idempotent_creation_retry_never_duplicates_or_reapplies_family_status(multi_app):
    main,admin=multi_app
    item=current(admin);key=str(uuid4())
    data=payload(item,status='completed',outcome='liked',completed_at='2026-10-24T18:00:00+03:00',request_id=key)
    first=admin.post('/api/listings/123/visits',json=data).json()
    before=first['visits']
    admin.post('/api/listings/123/review',json={'status':'new'})
    for _ in range(3):
        replay=admin.post('/api/listings/123/visits',json=data)
        assert replay.status_code==200
        assert replay.json()['visits']==before
        assert replay.json()['visits_version']==first['visits_version']
        assert replay.json()['review_status']=='new'
    assert admin.post('/api/listings/123/visits',json={**data,'comment':'Другой запрос'}).status_code==409
    other=current(admin,'456')
    assert admin.post('/api/listings/456/visits',json={**data,'property_id':other['property_id']}).status_code==409
    assert len(current(admin)['visits'])==1 and current(admin,'456')['visits']==[]


def test_simultaneous_duplicate_create_is_one_visit(multi_app):
    main,admin=multi_app
    data=payload(current(admin),request_id=str(uuid4()))
    barrier=Barrier(2)
    def save(_):
        barrier.wait(timeout=5)
        return admin.post('/api/listings/123/visits',json=data)
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses=list(pool.map(save,range(2)))
    assert [r.status_code for r in responses]==[200,200]
    assert len(current(admin)['visits'])==1
    assert current(admin)['visits_version']==1
