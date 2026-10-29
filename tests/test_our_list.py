"""Two authenticated pages and persistent property membership."""
import importlib
import io
import re
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.helpers import sign_in
from tests.test_multiuser import multi_app, reviewer_client
from tests.test_phase3 import sample
from tests.test_viewings import current, create

URL='/api/listings/123/selection'


def choose(client,item,selected=True,url=URL):
    return client.patch(url,json={'property_id':item['property_id'],
        'selection_version':item['selection_version'],'selected':selected})


def test_two_urls_share_session_and_expose_only_their_tools(multi_app):
    main,admin=multi_app
    for path,mode,count in [('/','all',24),('/ours','ours',28)]:
        response=admin.get(path)
        assert response.status_code==200 and response.headers['cache-control']=='no-store'
        html=response.text
        assert 'data-page="'+mode+'"' in html and '@@' not in html
        assert 'href="/"' in html and 'href="/ours"' in html
        assert len(re.findall(r'<th(?:\s|>)',html))==count
        assert re.search(r'/static/app.js\?v=[a-f0-9]{16}',html)
        assert 'visit-preparation' not in html
    all_html=admin.get('/').text;ours=admin.get('/ours').text
    for marker in ('visit-dialog','viewed-section','visit-agenda','compare-dialog','data-sort="visit_date"','data-review-filter="viewed"'):
        assert marker not in all_html
    assert 'id="filters"' in all_html and 'id="rejected-section"' in all_html
    assert 'id="filters"' not in ours and 'id="rejected-section"' not in ours
    for name in ('К просмотру','Оставляем в запасе','Проверка документов','Не понравилось','Риелтор / юрист','Этап отбора'):
        assert name in ours
    with TestClient(main.app,base_url='https://testserver') as stranger:
        for path in ('/','/ours'):
            response=stranger.get(path,follow_redirects=False)
            assert response.status_code==303 and response.headers['location']=='/login'


def test_membership_is_independent_preserves_history_and_survives_restart(multi_app):
    main,admin=multi_app
    assert not current(admin)['in_our_list']
    admin.post('/api/listings/123/review',json={'status':'liked'})
    assert not current(admin)['in_our_list']
    item=choose(admin,current(admin)).json()
    assert item['in_our_list'] and item['selection_version']==1
    same=choose(admin,item).json()
    assert same['selection_version']==1
    item=create(admin,status='completed',outcome='rejected',comment='Сохраняем историю',preparation='Старая заметка')
    assert item['in_our_list'] and item['viewed_group']=='rejected'
    admin.post('/api/listings/123/realtor-review',json={'status':'rejected','comment':'Документы'})
    item=current(admin);history=item['visits'];version=item['visits_version']
    removed=choose(admin,item,False).json()
    assert not removed['in_our_list'] and removed['review_status']=='rejected'
    assert removed['realtor_status']=='rejected' and removed['visits']==history
    assert removed['visits_version']==version
    from app.migrations import migrate_listings
    migrate_listings(main.database);migrate_listings(main.database)
    importlib.reload(main)
    assert not current(admin)['in_our_list']  # no repeated automatic selection
    assert current(admin)['visits']==history
    restored=choose(admin,current(admin)).json()
    assert restored['in_our_list'] and restored['visits']==history
    book=load_workbook(io.BytesIO(admin.get('/api/export.xlsx').content))
    sheet=book['Объявления']
    column=next(c.column for c in sheet[1] if c.value=='В нашем списке')
    row=next(row for row in range(2,sheet.max_row+1) if sheet.cell(row,1).value=='123')
    assert sheet.cell(row,column).value=='Да'
    assert book['Просмотры'].cell(2,8).value=='Старая заметка'


def test_membership_permissions_csrf_conflicts_and_concurrent_requests(multi_app):
    main,admin=multi_app
    _,buyer=reviewer_client(main,admin)
    old=current(admin)
    assert choose(buyer,old).status_code==200
    assert choose(admin,old,False).status_code==409
    assert choose(admin,{**current(admin),'property_id':current(admin,'456')['property_id']}).status_code==409
    assert choose(admin,current(admin),url='/api/listings/missing/selection').status_code==404
    assert admin.post('/api/users',json={'username':'expert','display_name':'Эксперт','password':'12345678','role':'realtor'}).status_code==201
    with TestClient(main.app,base_url='https://testserver') as expert:
        sign_in(expert,'expert','12345678')
        assert choose(expert,current(admin),False).status_code==403
    with TestClient(main.app,base_url='https://testserver') as no_csrf:
        sign_in(no_csrf);no_csrf.headers.pop('X-CSRF-Token')
        assert choose(no_csrf,current(admin),False).status_code==403
    with TestClient(main.app,base_url='https://testserver') as stranger:
        assert choose(stranger,current(admin),False).status_code in (401,403)
    item=current(admin);barrier=Barrier(2)
    def run(_):
        barrier.wait(timeout=5)
        return choose(admin,item,False).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(run,range(2)))==[200,409]
    assert not current(admin)['in_our_list']
    assert buyer.post('/api/listings/123/realtor-review',json={'status':'approved'}).status_code==403


@pytest.mark.parametrize('changes',[{'selected':1},{'selected':'true'},{'selected':None},
    {'selection_version':-1},{'review_status':'liked'}])
def test_membership_strict_input(multi_app,changes):
    main,admin=multi_app;item=current(admin)
    data={'property_id':item['property_id'],'selection_version':item['selection_version'],'selected':True,**changes}
    assert admin.patch(URL,json=data).status_code==422
    assert not current(admin)['in_our_list']


def test_one_time_migration_selects_any_existing_visit_only_and_preserves_vendor_data(multi_app):
    main,admin=multi_app
    with main.database() as conn:
        main.save_items(conn,[sample('700'),sample('701'),sample('702'),sample('703')],'2026-10-01T10:00:00Z')
    planned=create(admin,current(admin,'700'))
    completed=create(admin,current(admin,'701'),status='completed',outcome='thinking',preparation='Скрытая старая заметка')
    cancelled=create(admin,current(admin,'702'),status='cancelled',scheduled_at='2026-10-01T10:00:00Z')
    admin.post('/api/listings/703/review',json={'status':'liked'})
    with main.database() as conn:
        conn.execute('ALTER TABLE properties DROP COLUMN in_our_list')
        conn.execute('ALTER TABLE properties DROP COLUMN selection_version')
        conn.execute("DELETE FROM schema_migrations WHERE name='our_list_initial_selection_v1'")
        conn.execute('ALTER TABLE properties ADD COLUMN vendor_selection BLOB')
        conn.execute('UPDATE properties SET vendor_selection=?',(bytes([0,255]),))
        conn.execute('CREATE INDEX vendor_selection_idx ON properties(vendor_selection)')
        conn.execute('CREATE TRIGGER vendor_selection_trigger AFTER UPDATE OF vendor_selection ON properties BEGIN SELECT 1; END')
        before=[dict(r) for r in conn.execute('SELECT * FROM properties ORDER BY id')]
        visits=[dict(r) for r in conn.execute('SELECT * FROM property_visits ORDER BY id')]
        listings=[dict(r) for r in conn.execute('SELECT * FROM listings ORDER BY id')]
        schema=[tuple(r) for r in conn.execute("SELECT name,sql FROM sqlite_master WHERE name LIKE 'vendor_selection%' ORDER BY name")]
    from app.migrations import migrate_listings
    migrate_listings(main.database);migrate_listings(main.database)
    with main.database() as conn:
        for old,row in zip(before,conn.execute('SELECT * FROM properties ORDER BY id')):
            assert {k:row[k] for k in old}==old
        assert [dict(r) for r in conn.execute('SELECT * FROM property_visits ORDER BY id')]==visits
        assert [dict(r) for r in conn.execute('SELECT * FROM listings ORDER BY id')]==listings
        assert [tuple(r) for r in conn.execute("SELECT name,sql FROM sqlite_master WHERE name LIKE 'vendor_selection%' ORDER BY name")]==schema
        assert conn.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert conn.execute('PRAGMA foreign_key_check').fetchone() is None
    for key in ('700','701','702'):
        assert current(admin,key)['in_our_list'] and current(admin,key)['selection_version']==1
    assert not current(admin,'703')['in_our_list']
    assert current(admin,'701')['viewed_group'] is None
    assert current(admin,'701')['visits'][0]['outcome']=='thinking'
    removed=choose(admin,current(admin,'700'),False,url='/api/listings/700/selection')
    assert removed.status_code==200
    migrate_listings(main.database)
    assert not current(admin,'700')['in_our_list']
    assert admin.get('/api/auth/me').status_code==200


@pytest.mark.parametrize('selected_source',[False,True])
def test_merge_union_relisting_and_import_keep_membership(multi_app,selected_source):
    main,admin=multi_app
    with main.database() as conn:
        main.save_items(conn,[sample('700'),sample('701')],'2026-10-01T10:00:00Z')
    source=current(admin,'700');target=current(admin,'701')
    selected=source if selected_source else target
    selected=choose(admin,selected,url='/api/listings/'+selected['id']+'/selection').json()
    from app.property_match import merge_properties
    with main.database() as conn:
        conn.execute('BEGIN IMMEDIATE')
        merge_properties(conn,source['property_id'],target['property_id'],'2026-10-26T10:00:00Z')
    merged=current(admin,'700')
    assert merged['in_our_list'] and merged['selection_version']>selected['selection_version']
    assert choose(admin,source,False,url='/api/listings/700/selection').status_code==409
    with main.database() as conn:
        main.save_items(conn,[sample(merged['id'],price=12000000)],'2026-10-27T10:00:00Z',imported=True)
        main.save_items(conn,[sample('702',price=11900000)],'2026-10-28T10:00:00Z')
    final=current(admin,'700')
    assert final['in_our_list'] and '702' in final['listing_ids']
    assert len([i for i in admin.get('/api/state').json()['listings'] if '700' in i['listing_ids']])==1
    assert len(final['price_history'])>=3
