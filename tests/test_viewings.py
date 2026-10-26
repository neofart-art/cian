import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.helpers import sign_in
from tests.test_multiuser import multi_app, reviewer_client
from tests.test_phase3 import sample


def current(admin, listing='123'):
    return next(i for i in admin.get('/api/state').json()['listings'] if str(listing) in i['listing_ids'])


def payload(item, **changes):
    return {'property_id':item['property_id'],'status':'planned','scheduled_at':'2026-11-01T18:30:00+03:00',
            'preparation':'Спросить про окна',**changes}


def create(admin, item=None, **changes):
    item = item or current(admin)
    response = admin.post('/api/listings/'+item['id']+'/visits',json=payload(item,**changes))
    assert response.status_code == 200, response.text
    return response.json()


def test_planning_rescheduling_cancelling_preserves_decisions_and_history(multi_app):
    main, admin = multi_app
    admin.post('/api/listings/123/review',json={'status':'liked'})
    admin.post('/api/listings/123/realtor-review',json={'status':'rejected','comment':'Документы'})
    item = create(admin)
    assert item['review_status'] == 'liked' and item['realtor_status'] == 'rejected'
    visit = item['visits'][0]
    assert visit['scheduled_at'] == '2026-11-01T15:30:00+00:00'
    assert visit['updated_by_name'] and visit['version'] == 1
    for state in ('planned','cancelled'):
        response = admin.patch('/api/listings/123/visits/'+str(visit['id']),json=payload(item,
                    version=visit['version'],status=state,scheduled_at='2026-11-02T19:00:00+03:00'))
        assert response.status_code == 200
        item = response.json();visit=item['visits'][0]
        assert item['review_status']=='liked' and item['realtor_status']=='rejected'
    assert len(item['visits']) == 1 and item['visits_version'] == 3
    from app.migrations import migrate_listings
    migrate_listings(main.database);migrate_listings(main.database)
    assert current(admin)['visits'] == item['visits']
    with main.database() as conn:
        assert conn.execute('PRAGMA foreign_key_check').fetchone() is None


@pytest.mark.parametrize('outcome,status',[('liked','liked'),('rejected','rejected'),('thinking','viewed')])
def test_result_atomic_status_and_note_edit_does_not_undo_manual_decision(multi_app,outcome,status):
    main, admin=multi_app
    item=create(admin,status='completed',outcome=outcome,scheduled_at=None,comment='Впечатление')
    assert item['review_status']==status
    assert item['review_updated_by']==item['visits'][0]['updated_by']
    admin.post('/api/listings/123/review',json={'status':'new'})
    v=item['visits'][0]
    response=admin.patch('/api/listings/123/visits/'+str(v['id']),json=payload(item,
        version=v['version'],status='completed',outcome=outcome,comment='Уточнение'))
    assert response.json()['review_status']=='new'
    # Scheduling/cancelling a repeat cannot erase the first outcome.
    repeat=create(admin)
    assert len(repeat['visits'])==2 and repeat['review_status']=='new'
    assert any(x['outcome']==outcome for x in repeat['visits'])
    book=load_workbook(io.BytesIO(admin.get('/api/export.xlsx').content))
    assert book['Просмотры'].max_row==3


def test_stale_edits_bad_data_and_roles(multi_app):
    main,admin=multi_app
    _,reviewer=reviewer_client(main,admin)
    item=create(reviewer)
    v=item['visits'][0];url='/api/listings/123/visits/'+str(v['id'])
    edit=payload(item,version=v['version'],status='completed',outcome='rejected',comment='=unsafe')
    assert admin.patch(url,json=edit).status_code==200
    assert reviewer.patch(url,json=edit).status_code==409
    assert current(admin)['review_status']=='rejected'
    assert admin.patch(url,json=payload(item,version=2,status='cancelled')).status_code==422
    for changes in ({'scheduled_at':None},{'scheduled_at':'2026-11-01T18:30'},
                    {'status':'completed'},{'status':'planned','outcome':'liked'}, {'preparation':'a\x00b'},
                    {'comment':'x'*2001}):
        assert admin.post('/api/listings/123/visits',json=payload(item,**changes)).status_code==422
    assert admin.post('/api/listings/123/visits',json=payload(item,property_id=99999)).status_code==409
    admin.post('/api/users',json={'username':'lawyer','display_name':'Юрист','password':'password-123','role':'realtor'})
    with TestClient(main.app,base_url='https://testserver') as expert:
        sign_in(expert,'lawyer','password-123')
        assert expert.get('/api/state').json()['listings'][0]['visits'] is not None
        assert expert.post('/api/listings/123/visits',json=payload(item)).status_code==403
        assert expert.patch(url,json=edit).status_code==403
    with TestClient(main.app,base_url='https://testserver') as stranger:
        assert stranger.get('/api/state').status_code==401
        assert stranger.post('/api/listings/123/visits',json=payload(item)).status_code in (401,403)
    with TestClient(main.app,base_url='https://testserver') as no_csrf:
        sign_in(no_csrf);no_csrf.headers.pop('X-CSRF-Token')
        assert no_csrf.post('/api/listings/123/visits',json=payload(item)).status_code==403
    assert len(current(admin)['visits'])==1
    book=load_workbook(io.BytesIO(admin.get('/api/export.xlsx').content))
    assert book['Просмотры'].cell(2,9).value=="'=unsafe"


def test_relisting_merge_and_import_keep_viewings(multi_app):
    main,admin=multi_app
    with main.database() as conn:
        main.save_items(conn,[sample('700'),sample('701')],'2026-10-01T10:00:00+00:00')
    a=create(admin,current(admin,'700'),status='completed',outcome='liked')
    b=create(admin,current(admin,'701'))
    from app.property_match import merge_properties
    with main.database() as conn:
        conn.execute('BEGIN IMMEDIATE')
        merge_properties(conn,a['property_id'],b['property_id'],'2026-10-26T10:00:00+00:00')
        main.save_items(conn,[sample('702')],'2026-11-01T10:00:00+00:00')
    merged=current(admin,'700')
    assert len(merged['visits'])==2 and merged['visits_version']>max(a['visits_version'],b['visits_version'])
    # A stale source-property editor must not silently write into the merged target.
    assert admin.post('/api/listings/700/visits',json=payload(a)).status_code==409
    assert merged['review_status']=='liked'
    with main.database() as conn:
        main.save_items(conn,[sample(merged['id'],price=12000000)],'2026-11-02T10:00:00+00:00',imported=True)
    assert current(admin,'700')['visits']==merged['visits']


def test_older_viewing_correction_does_not_override_newer_result(multi_app):
    main,admin=multi_app
    first=create(admin,status='completed',outcome='liked')
    old=first['visits'][0]
    create(admin,status='completed',outcome='rejected')
    response=admin.patch('/api/listings/123/visits/'+str(old['id']),json=payload(first,
                         version=old['version'],status='completed',outcome='thinking'))
    assert response.status_code==200 and response.json()['review_status']=='rejected'
