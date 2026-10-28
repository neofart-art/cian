"""Property-level post-visit classification, additive migration and permissions."""
import importlib
import io
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.helpers import sign_in
from tests.test_multiuser import multi_app, reviewer_client
from tests.test_phase3 import sample
from tests.test_viewings import current, create, payload

URL='/api/listings/123/viewing-classification'


def stage(client,item,ready=True,url=URL):
    return client.patch(url,json={'property_id':item['property_id'],
        'visits_version':item['visits_version'],'documents_ready':ready})


def edit(client,item,visit,**changes):
    response=client.patch('/api/listings/'+item['id']+'/visits/'+str(visit['id']),json=payload(item,
        **{'version':visit['version'],'status':'completed','outcome':visit['outcome'],
           'completed_at':visit['completed_at'],**changes}))
    assert response.status_code==200,response.text
    return response.json()


def test_result_not_family_or_realtor_status_controls_group(multi_app):
    main,admin=multi_app
    assert current(admin)['viewed_group'] is None
    admin.post('/api/listings/123/review',json={'status':'liked'})
    assert current(admin)['viewed_group'] is None
    admin.post('/api/listings/123/realtor-review',json={'status':'rejected','comment':'Проверить документы'})
    assert current(admin)['viewed_group'] is None
    item=create(admin,status='completed',outcome='liked',completed_at='2026-10-20T10:00:00Z')
    assert item['viewed_group']=='liked'
    assert item['realtor_status']=='rejected' and item['realtor_comment']=='Проверить документы'
    item=stage(admin,item).json()
    assert item['viewed_group']=='documents' and item['review_status']=='liked'
    admin.post('/api/listings/123/review',json={'status':'rejected'})
    assert current(admin)['viewed_group']=='documents'
    item=stage(admin,current(admin),False).json()
    assert item['viewed_group']=='liked' and item['review_status']=='rejected'
    assert item['realtor_status']=='rejected'


def test_preparation_survives_notes_dates_repeat_and_reload(multi_app):
    main,admin=multi_app
    item=create(admin,status='completed',outcome='liked',completed_at='2026-10-20T10:00:00Z')
    item=stage(admin,item).json()
    v=item['visits'][0]
    item=edit(admin,item,v,comment='Уточнили после поездки',preparation='Вопрос до поездки')
    assert item['viewed_group']=='documents'
    v=item['visits'][0]
    response=admin.patch('/api/listings/123/visits/date',json={'property_id':item['property_id'],
        'visits_version':item['visits_version'],'visit_id':v['id'],'date':'2026-10-21'})
    assert response.status_code==200 and response.json()['viewed_group']=='documents'
    # A second positive result and a new appointment do not erase preparation.
    item=create(admin,status='completed',outcome='liked',completed_at='2026-10-22T10:00:00Z')
    assert item['viewed_group']=='documents'
    item=create(admin)
    assert item['viewed_group']=='documents' and len(item['visits'])==3
    importlib.reload(main)
    assert current(admin)['viewed_group']=='documents'
    assert current(admin)['visits']==item['visits']


def test_latest_actual_result_and_negative_correction_reset_only_latest(multi_app):
    main,admin=multi_app
    item=create(admin,status='completed',outcome='rejected',completed_at='2026-10-20T10:00:00Z')
    old=item['visits'][0]
    item=create(admin,status='completed',outcome='liked',completed_at='2026-10-23T10:00:00Z')
    item=stage(admin,item).json()
    # Older result correction cannot override the later positive or its group.
    item=edit(admin,item,old,**{'outcome':'liked'})
    assert item['viewed_group']=='documents' and item['review_status']=='liked'
    item=create(admin,status='completed',outcome='rejected',completed_at='2026-10-24T10:00:00Z',
                rejection_reasons=['noise'])
    assert item['viewed_group']=='rejected' and item['review_status']=='rejected'
    assert stage(admin,item).status_code==409
    item=edit(admin,item,item['visits'][0],**{'outcome':'liked'})
    assert item['viewed_group']=='liked' and not item['viewing_documents_ready']
    assert item['review_status']=='liked' and len(item['visits'])==3
    assert all(v['status']=='completed' for v in item['visits'])


def test_actual_date_reordering_changes_result_group_but_keeps_preparation(multi_app):
    main,admin=multi_app
    item=create(admin,status='completed',outcome='rejected',completed_at='2026-10-20T10:00:00Z')
    item=create(admin,status='completed',outcome='liked',completed_at='2026-10-23T10:00:00Z')
    item=stage(admin,item).json()
    positive=item['visits'][0]
    item=edit(admin,item,positive,**{'completed_at':'2026-10-19T10:00:00Z'})
    assert item['viewed_group']=='rejected' and item['viewing_documents_ready']
    item=edit(admin,item,item['visits'][0],**{'completed_at':'2026-10-25T10:00:00Z'})
    assert item['viewed_group']=='documents'


def test_classification_roles_csrf_stale_property_and_concurrent_edits(multi_app):
    main,admin=multi_app
    _,buyer=reviewer_client(main,admin)
    item=create(admin,status='completed',outcome='liked')
    assert stage(buyer,item).status_code==200
    assert stage(admin,item,False).status_code==409
    item=current(admin)
    assert stage(admin,item,False).status_code==200
    assert admin.post('/api/users',json={'username':'expert','display_name':'Эксперт','password':'12345678','role':'realtor'}).status_code==201
    with TestClient(main.app,base_url='https://testserver') as expert:
        sign_in(expert,'expert','12345678')
        assert stage(expert,current(admin)).status_code==403
    with TestClient(main.app,base_url='https://testserver') as no_csrf:
        sign_in(no_csrf);no_csrf.headers.pop('X-CSRF-Token')
        assert stage(no_csrf,current(admin)).status_code==403
    with TestClient(main.app,base_url='https://testserver') as stranger:
        assert stage(stranger,current(admin)).status_code in (401,403)
    item=current(admin)
    bad={**item,'property_id':current(admin,'456')['property_id']}
    assert stage(admin,bad).status_code==409
    assert stage(admin,item,url='/api/listings/missing/viewing-classification').status_code==404
    barrier=Barrier(2)
    def save(_):
        barrier.wait(timeout=5)
        return stage(admin,item).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(save,range(2)))==[200,409]
    assert current(admin)['viewed_group']=='documents'
    assert buyer.post('/api/listings/123/realtor-review',json={'status':'approved','comment':''}).status_code==403


@pytest.mark.parametrize('changes',[{'documents_ready':'true'},{'documents_ready':1},
    {'documents_ready':None},{'unexpected':'x'},{'visits_version':-1}])
def test_classification_strict_input(multi_app,changes):
    main,admin=multi_app
    item=create(admin,status='completed',outcome='liked')
    data={'property_id':item['property_id'],'visits_version':item['visits_version'],'documents_ready':True,**changes}
    assert admin.patch(URL,json=data).status_code==422
    assert current(admin)['viewed_group']=='liked'


def test_legacy_results_additive_idempotent_migration_no_reclassification(multi_app):
    main,admin=multi_app
    create(admin,status='completed',outcome='thinking',comment='Не решили')
    create(admin,current(admin,'456'),status='completed',outcome='liked')
    with main.database() as conn:
        conn.execute('ALTER TABLE properties DROP COLUMN viewing_documents_ready')
        conn.execute('ALTER TABLE properties ADD COLUMN vendor_blob BLOB')
        conn.execute('UPDATE properties SET vendor_blob=?',(bytes([0,255]),))
        conn.execute('CREATE INDEX vendor_property_idx ON properties(vendor_blob)')
        conn.execute('CREATE TRIGGER vendor_property_trigger AFTER UPDATE OF vendor_blob ON properties BEGIN SELECT 1; END')
        before=[dict(r) for r in conn.execute('SELECT * FROM properties ORDER BY id')]
        visits=[dict(r) for r in conn.execute('SELECT * FROM property_visits ORDER BY id')]
        listings=[dict(r) for r in conn.execute('SELECT * FROM listings ORDER BY id')]
        schema=[tuple(r) for r in conn.execute("SELECT name,sql FROM sqlite_master WHERE name LIKE 'vendor_%' ORDER BY name")]
    from app.migrations import migrate_listings
    migrate_listings(main.database);migrate_listings(main.database)
    with main.database() as conn:
        for old,row in zip(before,conn.execute('SELECT * FROM properties ORDER BY id')):
            assert {k:row[k] for k in old}==old
            assert row['viewing_documents_ready']==0
        assert [dict(r) for r in conn.execute('SELECT * FROM property_visits ORDER BY id')]==visits
        assert [dict(r) for r in conn.execute('SELECT * FROM listings ORDER BY id')]==listings
        assert [tuple(r) for r in conn.execute("SELECT name,sql FROM sqlite_master WHERE name LIKE 'vendor_%' ORDER BY name")]==schema
        assert conn.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert conn.execute('PRAGMA foreign_key_check').fetchone() is None
    assert current(admin)['viewed_group'] is None
    assert current(admin)['visits'][0]['outcome']=='thinking'
    assert stage(admin,current(admin)).status_code==409
    assert current(admin,'456')['viewed_group']=='liked'
    assert admin.get('/api/auth/me').status_code==200


@pytest.mark.parametrize('reverse',[False,True])
def test_merge_import_and_republication_preserve_latest_property_classification(multi_app,reverse):
    main,admin=multi_app
    with main.database() as conn:
        main.save_items(conn,[sample('700'),sample('701')],'2026-10-01T10:00:00+00:00')
    a=create(admin,current(admin,'700'),status='completed',outcome='liked',completed_at='2026-10-24T10:00:00Z')
    a=stage(admin,a,url='/api/listings/700/viewing-classification').json()
    b=create(admin,current(admin,'701'),status='completed',outcome='rejected',completed_at='2026-10-20T10:00:00Z')
    from app.property_match import merge_properties
    source,target=(a,b) if not reverse else (b,a)
    with main.database() as conn:
        conn.execute('BEGIN IMMEDIATE')
        merge_properties(conn,source['property_id'],target['property_id'],'2026-10-26T10:00:00+00:00')
    item=current(admin,'700')
    assert item['viewed_group']=='documents' and len(item['visits'])==2
    assert item['visits_version']>max(a['visits_version'],b['visits_version'])
    assert stage(admin,a,url='/api/listings/700/viewing-classification').status_code==409
    with main.database() as conn:
        main.save_items(conn,[sample(item['id'],price=12500000)],'2026-10-27T10:00:00+00:00',imported=True)
        main.save_items(conn,[sample('702',price=12400000)],'2026-10-28T10:00:00+00:00')
    item=current(admin,'700')
    assert '702' in item['listing_ids'] and item['viewed_group']=='documents'
    assert len(item['visits'])==2 and len(item['price_history'])>=3
    book=load_workbook(io.BytesIO(admin.get('/api/export.xlsx').content))
    sheet=book['Объявления']
    column=next(c.column for c in sheet[1] if c.value=='Группа после просмотра')
    assert any(sheet.cell(row,column).value=='Понравилось — готовимся к проверке документов' for row in range(2,sheet.max_row+1))
    assert book['Просмотры'].max_row==3


def test_correcting_negative_with_new_actual_date_starts_without_documents(multi_app):
    main,admin=multi_app
    item=create(admin,status='completed',outcome='rejected',completed_at='2026-10-20T10:00:00Z')
    old=item['visits'][0]
    item=create(admin,status='completed',outcome='liked',completed_at='2026-10-22T10:00:00Z')
    item=stage(admin,item).json()
    item=edit(admin,item,old,**{'outcome':'liked','completed_at':'2026-10-23T10:00:00Z'})
    assert item['viewed_group']=='liked' and not item['viewing_documents_ready']
    assert len(item['visits'])==2
